"""Small desktop interface for sending a prompt once per populated Excel cell."""
from __future__ import annotations

import asyncio
import csv
import queue
import re
import threading
from datetime import datetime
from pathlib import Path
from tkinter import BOTH, END, LEFT, RIGHT, TOP, filedialog, messagebox
import tkinter as tk
from tkinter import ttk

from .router import ModelUnavailableError, Providers


def read_workbook(path: Path) -> list[tuple[str, str]]:
    """Return populated cells in worksheet order, then row/column order."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - dependency is declared
        raise RuntimeError("Excel support requires openpyxl. Reinstall smart_AI with pip install -e .") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    values: list[tuple[str, str]] = []
    try:
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is not None and str(cell.value).strip():
                        values.append((f"{sheet.title}!{cell.coordinate}", str(cell.value)))
    finally:
        workbook.close()
    return values


def read_csv(path: Path) -> list[tuple[str, str]]:
    """Return populated CSV cells in row and column order."""
    values: list[tuple[str, str]] = []
    # utf-8-sig also handles the BOM commonly added by spreadsheet programs.
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        for row_number, row in enumerate(csv.reader(stream), 1):
            for column_number, value in enumerate(row, 1):
                if value.strip():
                    values.append((f"R{row_number}C{column_number}", value))
    return values


def read_cells(path: Path) -> list[tuple[str, str]]:
    if path.suffix.lower() == ".csv":
        return read_csv(path)
    return read_workbook(path)


class PromptQueueApp:
    def __init__(self, root: tk.Tk) -> None:
        self.root = root
        root.title("smart_AI Excel Prompt Runner")
        root.geometry("760x650")
        root.minsize(620, 520)
        self.workbook = tk.StringVar()
        self.output_dir = tk.StringVar(value=str(Path.cwd() / "results"))
        self.placeholder = tk.StringVar(value="##")
        self.delay = tk.StringVar(value="2")
        self.save_results = tk.BooleanVar(value=True)
        self.stop_event = threading.Event()
        self.events: queue.Queue = queue.Queue()
        self.running = False

        outer = ttk.Frame(root, padding=12)
        outer.pack(fill=BOTH, expand=True)
        ttk.Label(outer, text="Excel workbook").pack(anchor="w")
        row = ttk.Frame(outer); row.pack(fill="x", pady=(3, 10))
        ttk.Entry(row, textvariable=self.workbook).pack(side=LEFT, fill="x", expand=True)
        ttk.Button(row, text="Browse…", command=self.pick_workbook).pack(side=RIGHT, padx=(8, 0))
        ttk.Label(outer, text="Prompt template (the placeholder is replaced with each cell value)").pack(anchor="w")
        self.prompt = tk.Text(outer, height=10, wrap="word")
        self.prompt.pack(fill=BOTH, expand=True, pady=(3, 10))
        self.prompt.insert("1.0", "Please process this item: ##")
        options = ttk.Frame(outer); options.pack(fill="x", pady=(0, 8))
        ttk.Label(options, text="Placeholder").pack(side=LEFT)
        ttk.Entry(options, width=9, textvariable=self.placeholder).pack(side=LEFT, padx=(5, 16))
        ttk.Label(options, text="Delay between requests (seconds)").pack(side=LEFT)
        ttk.Entry(options, width=8, textvariable=self.delay).pack(side=LEFT, padx=5)
        ttk.Checkbutton(outer, text="Save each response as a separate .txt file", variable=self.save_results).pack(anchor="w")
        outrow = ttk.Frame(outer); outrow.pack(fill="x", pady=(5, 8))
        ttk.Label(outrow, text="Results folder").pack(side=LEFT)
        ttk.Entry(outrow, textvariable=self.output_dir).pack(side=LEFT, fill="x", expand=True, padx=8)
        ttk.Button(outrow, text="Browse…", command=self.pick_output).pack(side=RIGHT)
        buttons = ttk.Frame(outer); buttons.pack(fill="x", pady=(2, 8))
        self.start_button = ttk.Button(buttons, text="Start", command=self.start)
        self.start_button.pack(side=LEFT)
        self.stop_button = ttk.Button(buttons, text="Stop after current request", command=self.stop, state="disabled")
        self.stop_button.pack(side=LEFT, padx=8)
        self.status = ttk.Label(buttons, text="Ready")
        self.status.pack(side=RIGHT)
        self.log = tk.Text(outer, height=9, state="disabled", wrap="word")
        self.log.pack(fill=BOTH, expand=True)
        root.after(100, self.poll_events)

    def pick_workbook(self) -> None:
        name = filedialog.askopenfilename(filetypes=[("Spreadsheets", "*.xlsx *.xlsm *.csv"), ("Excel workbooks", "*.xlsx *.xlsm"), ("CSV files", "*.csv"), ("All files", "*.*")])
        if name: self.workbook.set(name)

    def pick_output(self) -> None:
        name = filedialog.askdirectory(initialdir=self.output_dir.get() or None)
        if name: self.output_dir.set(name)

    def append_log(self, text: str) -> None:
        self.log.configure(state="normal")
        self.log.insert(END, text + "\n")
        self.log.see(END)
        self.log.configure(state="disabled")

    def start(self) -> None:
        if self.running: return
        path = Path(self.workbook.get()).expanduser()
        template = self.prompt.get("1.0", END).strip()
        marker = self.placeholder.get()
        try: delay = float(self.delay.get())
        except ValueError:
            messagebox.showerror("Invalid delay", "Enter a delay of 0 seconds or more."); return
        if not path.is_file() or path.suffix.lower() not in {".xlsx", ".xlsm", ".csv"}:
            messagebox.showerror("Choose a spreadsheet", "Select an existing .xlsx, .xlsm, or .csv file."); return
        if not template or not marker or marker not in template:
            messagebox.showerror("Prompt placeholder", "The prompt must contain the placeholder text."); return
        if delay < 0:
            messagebox.showerror("Invalid delay", "Enter a delay of 0 seconds or more."); return
        try: items = read_cells(path)
        except Exception as exc:
            messagebox.showerror("Could not read workbook", str(exc)); return
        if not items:
            messagebox.showinfo("No cell content", "No populated cells were found in this workbook."); return
        folder = Path(self.output_dir.get()).expanduser()
        if self.save_results.get():
            try: folder.mkdir(parents=True, exist_ok=True)
            except OSError as exc:
                messagebox.showerror("Results folder", str(exc)); return
        self.running = True
        self.stop_event.clear()
        self.start_button.configure(state="disabled")
        self.stop_button.configure(state="normal")
        self.status.configure(text=f"Starting 0/{len(items)}")
        self.append_log(f"Loaded {len(items)} populated cells from {path.name}.")
        threading.Thread(target=self.run_queue, args=(items, template, marker, delay, folder, self.save_results.get()), daemon=True).start()

    def stop(self) -> None:
        self.stop_event.set()
        self.status.configure(text="Stopping after current request…")

    def _request_fallback_confirmation(self, error: ModelUnavailableError) -> bool:
        answer = threading.Event()
        decision: list[bool] = []
        self.events.put(("fallback_offer", str(error), answer, decision))
        while not answer.wait(0.1):
            if self.stop_event.is_set(): return False
        return decision[0] if decision else False

    def run_queue(self, items: list[tuple[str, str]], template: str, marker: str, delay: float, folder: Path, save: bool) -> None:
        done = 0
        run_id = datetime.now().strftime("%Y%m%d_%H%M%S")
        fallback_model: str | None = None
        try:
            for index, (cell, value) in enumerate(items, 1):
                if self.stop_event.is_set(): break
                prompt = template.replace(marker, value)
                async def send(model_override: str | None = None) -> dict:
                    providers = Providers()
                    try: return await providers.run({"input": prompt}, {}, model_override=model_override)
                    finally: await providers.close()
                try:
                    used_fallback = fallback_model is not None
                    result = asyncio.run(send(fallback_model))
                except ModelUnavailableError as exc:
                    if exc.requires_confirmation and fallback_model is None:
                        if not self._request_fallback_confirmation(exc):
                            self.stop_event.set()
                            break
                    async def try_fallbacks() -> dict:
                        providers = Providers()
                        try: return await providers.run_fallbacks({"input": prompt}, {}, exc.model)
                        finally: await providers.close()
                    result = asyncio.run(try_fallbacks())
                    fallback_model = result.get("model")
                    used_fallback = True
                output = result.get("output_text") or ""
                if not output: output = str(result.get("raw", {}))
                plan = result.get("plan") or {}
                active_tools = result.get("effective_request_options", {}).get("tools", [])
                planned_tools = plan.get("tools", [])
                search_enabled = "google_search" in planned_tools or any(
                    (tool.get("type") if isinstance(tool, dict) else tool) == "google_search"
                    for tool in active_tools
                )
                has_url = "http://" in prompt.lower() or "https://" in prompt.lower()
                route_note = f"model={result.get('model', 'unknown')}; Google Search={'on' if search_enabled else 'off'}"
                if used_fallback:
                    route_note += "; fallback bypassed Cloudflare planning"
                if has_url and not search_enabled:
                    route_note += "; URL not fetched (Search unavailable for this request)"
                if save:
                    safe_cell = re.sub(r"[^A-Za-z0-9_.-]+", "_", cell)
                    target = folder / f"{run_id}_{index:04d}_{safe_cell}.txt"
                    target.write_text(output, encoding="utf-8")
                done += 1
                self.events.put(("progress", done, len(items), cell, route_note, output[:180].replace("\n", " ")))
                if index < len(items) and self.stop_event.wait(delay): break
            self.events.put(("finished", done, len(items), self.stop_event.is_set()))
        except Exception as exc:
            self.events.put(("error", done, len(items), str(exc)))

    def poll_events(self) -> None:
        while True:
            try: event = self.events.get_nowait()
            except queue.Empty: break
            if event[0] == "progress":
                _, done, total, cell, route_note, excerpt = event
                self.status.configure(text=f"{done}/{total}")
                self.append_log(f"[{done}/{total}] {cell} [{route_note}]: {excerpt}")
            elif event[0] == "finished":
                _, done, total, stopped = event
                self.running = False
                self.start_button.configure(state="normal")
                self.stop_button.configure(state="disabled")
                self.status.configure(text=f"Stopped at {done}/{total}" if stopped else f"Complete: {done}/{total}")
                self.append_log(self.status.cget("text"))
            elif event[0] == "fallback_offer":
                _, detail, answer, decision = event
                accepted = messagebox.askyesno("Model unavailable", f"{detail}\n\nTry other locally eligible models and retry without another Cloudflare planning call? Fallback requests do not enable Google Search, so URL contents will not be fetched.")
                decision.append(accepted)
                answer.set()
            elif event[0] == "error":
                _, done, total, detail = event
                self.running = False
                self.start_button.configure(state="normal")
                self.stop_button.configure(state="disabled")
                self.status.configure(text=f"Failed at {done}/{total}")
                self.append_log(f"Error: {detail}")
                messagebox.showerror("Request failed", detail)
        self.root.after(100, self.poll_events)


def main() -> None:
    root = tk.Tk()
    PromptQueueApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
